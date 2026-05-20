"""
Generate NESMA requirements specification DOCX document.

Usage:
    python generate_docx.py --output <path.docx> --data <data.json>
"""

import argparse
import json
import sys
from pathlib import Path
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("请先安装 python-docx: pip install python-docx")
    sys.exit(1)


def set_cell_shading(cell, color):
    """Set cell background color."""
    shading_elm = OxmlElement("w:shd")
    shading_elm.set(qn("w:fill"), color)
    shading_elm.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_border(cell, **kwargs):
    """Set cell borders."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("start", "top", "end", "bottom"):
        edge_data = kwargs.get(edge)
        if edge_data:
            element = OxmlElement(f"w:{edge}")
            for attr in ["sz", "val", "color"]:
                if attr in edge_data:
                    element.set(qn(f"w:{attr}"), str(edge_data[attr]))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def make_border():
    """Standard thin border."""
    return {"sz": "4", "val": "single", "color": "000000"}


def add_styled_paragraph(doc, text, style="Normal", bold=False, size=10.5,
                         alignment=None, space_after=6):
    """Add a paragraph with styling."""
    p = doc.add_paragraph(text, style=style)
    if bold:
        for run in p.runs:
            run.bold = True
    if alignment is not None:
        p.alignment = alignment
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    return p


def build_project_info_table(doc, data):
    """Build project info table."""
    project_name = data.get("project_name", "")
    project_type = data.get("project_type", "新建项目")

    table = doc.add_table(rows=2, cols=3, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Row 1
    table.cell(0, 0).text = "项目名称："
    table.cell(0, 1).text = project_name
    table.cell(0, 2).text = project_name
    # Row 2
    table.cell(1, 0).text = "项目类型："
    table.cell(1, 1).text = project_type
    table.cell(1, 2).text = project_type

    # Merge cells
    table.cell(0, 1).merge(table.cell(0, 2))
    table.cell(1, 1).merge(table.cell(1, 2))

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    run.font.name = "微软雅黑"

    return table


def build_requirements_sections(doc, data):
    """Build functional requirements sections from function points."""
    functions = data.get("functions", [])

    # Group by level1
    modules = {}
    for f in functions:
        l1 = f.get("level1", "其他")
        l2 = f.get("level2", "")
        if l1 not in modules:
            modules[l1] = {}
        if l2 not in modules[l1]:
            modules[l1][l2] = []
        modules[l1][l2].append(f)

    for mod_name, sub_modules in modules.items():
        doc.add_heading(f"功能模块：{mod_name}", level=2)

        for sub_name, funcs in sub_modules.items():
            if sub_name:
                doc.add_heading(sub_name, level=3)

            for func in funcs:
                name = func.get("name", "")
                category = func.get("category", "")
                category_cn = {
                    "ILF": "内部逻辑文件(ILF)",
                    "EIF": "外部接口文件(EIF)",
                    "EI": "外部输入(EI)",
                    "EO": "外部输出(EO)",
                    "EQ": "外部查询(EQ)",
                }.get(category, category)

                p = doc.add_paragraph()
                run = p.add_run(f"功能点：{name}")
                run.bold = True
                run.font.size = Pt(10.5)

                p2 = doc.add_paragraph(f"类别：{category_cn}　　复用程度：{func.get('reuse', '低')}　　修改类型：{func.get('modify_type', '新增')}")
                p2.style = doc.styles["Normal"]
                for run in p2.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

                # Add detailed description placeholder
                p3 = doc.add_paragraph(
                    f"功能描述：请根据实际需求补充{name}的详细功能描述，"
                    f"包括输入数据、处理逻辑、输出结果等。"
                )
                for run in p3.runs:
                    run.font.size = Pt(10.5)


def generate_docx(output_path, data):
    """Generate complete requirements specification DOCX."""
    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "微软雅黑"
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    # === Title Page ===
    title = doc.add_heading(f"{data.get('project_name', '')}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_heading("需求规格说明书", level=1)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    version_para = doc.add_paragraph(f"阶段：立项阶段")
    version_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # === Project Info ===
    build_project_info_table(doc, data)
    doc.add_paragraph()

    # === Background ===
    doc.add_heading("1. 背景", level=1)
    doc.add_heading("1.1 基本信息", level=2)
    doc.add_paragraph("关于本项目的背景信息、立项依据、建设必要性等内容。")
    doc.add_heading("1.2 目标", level=2)
    doc.add_paragraph(
        f"本文档为《{data.get('project_name', '')}》的需求规格说明书（立项阶段），"
        f"主要目的是明确项目的功能需求范围和建设内容，作为软件工作量评估的重要依据。"
    )
    doc.add_heading("1.3 适用范围", level=2)
    doc.add_paragraph(
        "本文档适用于建设单位、承建单位、测试单位等项目相关方，"
        "作为项目设计、开发和验收的参考依据。"
    )
    doc.add_heading("1.4 参考文献", level=2)
    refs = [
        "GB/T 42588-2023 系统与软件工程 功能规模测量 NESMA方法",
        "中国软件行业基准数据报告（SSM-BK-202409）",
        "GB/T 36964-2018 软件工程 软件开发成本度量规范",
    ]
    for r in refs:
        doc.add_paragraph(r, style="List Bullet")

    # === Functional Requirements ===
    doc.add_heading("2. 功能需求", level=1)
    doc.add_paragraph(
        f"本项目共识别 {len(data.get('functions', []))} 个功能点，"
        f"采用 NESMA {data.get('method', '估算法')}进行功能规模度量。"
    )
    build_requirements_sections(doc, data)

    # === Non-functional Requirements ===
    doc.add_heading("3. 非功能需求", level=1)

    doc.add_heading("3.1 性能需求", level=2)
    doc.add_paragraph("系统应满足业务场景的性能要求，包括响应时间、并发用户数等。")
    doc.add_heading("3.2 可靠性需求", level=2)
    doc.add_paragraph("系统应提供可靠的服务保障，确定可接受的故障恢复时间。")
    doc.add_heading("3.3 安全性需求", level=2)
    doc.add_paragraph("系统应符合信息安全等级保护要求，做好访问控制和数据保护。")

    # === Architecture Overview ===
    doc.add_heading("4. 架构概述", level=1)
    doc.add_paragraph(
        f"本项目采用{data.get('project_info', {}).get('dev_language', '主流')}技术栈开发，"
        f"业务领域归属{data.get('project_info', {}).get('business_domain', '通用')}。"
        "具体架构设计详见后续详细设计文档。"
    )

    # === Appendix ===
    doc.add_heading("5. 附录", level=1)
    doc.add_paragraph("附录A：功能点拆分表（详见配套 Excel 文件）")
    doc.add_paragraph(f"附录B：编制日期 {datetime.now().strftime('%Y年%m月%d日')}")

    # Save
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"DOCX document saved to: {output_path}")
    print(f"  Sections: {len(doc.sections)}")
    print(f"  Paragraphs: {len(doc.paragraphs)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate NESMA requirements specification")
    parser.add_argument("--output", required=True, help="Output .docx file path")
    parser.add_argument("--data", required=True, help="JSON data file path")
    parser.add_argument("--project-name", help="Project name (optional)")
    args = parser.parse_args()

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    if args.project_name:
        data["project_name"] = args.project_name

    generate_docx(args.output, data)
